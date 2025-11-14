from docx import Document
from typing import List

def read_docx(path: str) -> str:
    doc = Document(path)
    texts = []
    for p in doc.paragraphs:
        texts.append(p.text)
    return "\n".join(texts)

def write_docx(path: str, content: str | List[str]):
    doc = Document()
    if isinstance(content, list):
        for p in content:
            doc.add_paragraph(p)
    else:
        doc.add_paragraph(content)
    doc.save(path)
    return path